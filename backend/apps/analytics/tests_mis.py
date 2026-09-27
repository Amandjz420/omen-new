"""The MIS engine and the money reports."""
import datetime
import io
from decimal import Decimal

from django.contrib.auth.models import User
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from target_core import models as m
from apps.migration_pipeline.services import source_key
from apps.valuations.test_forms import Form, user_with_role


class MisFixture(Form):
    """Shared setup and helpers; holds no tests, so none run twice."""

    def setUp(self):
        self.build()
        self.client.force_login(User.objects.create_user('staff', password='x', is_staff=True))
        self.today = timezone.localdate()

    def invoice(self, billed, remaining=0, days_ago=10, status='Submit', assigned=None, report=True, cancelled=False, **case):
        c = self.lead_with_case(assigned=assigned)
        for k, v in case.items():
            setattr(c, k, v)
        c.save()
        v = m.Valuation.objects.create(case=c, legacy_status=status, report_number=f'R-{m.Valuation.objects.count() + 1}')
        if report:
            m.BankReport.objects.create(valuation=v, revision=0)
        for head in range(6):                                              # six charge lines, as legacy invoices carry
            m.ValuationCharge.objects.create(valuation=v, amount=Decimal(billed) / 6)
        return m.Invoice.objects.create(valuation=v, origin='legacy_invoice', number=f'I-{m.Invoice.objects.count() + 1}',
                                        issued_date=self.today - datetime.timedelta(days=days_ago), billed_amount=Decimal(billed),
                                        legacy_paid_amount=Decimal(billed) - Decimal(remaining), legacy_remaining_amount=Decimal(remaining),
                                        cancelled_at=timezone.now() if cancelled else None)

    def get(self, key, **params):
        return self.client.get(f'/api/mis/{key}/', params)

class MisTests(MisFixture, TestCase):
    def test_money_is_counted_once_per_invoice_whatever_its_charge_lines(self):
        """Regression: the legacy query summed across six charge lines per invoice."""
        self.invoice('1180', remaining='680')
        t = self.get('invoices').json()['totals']
        self.assertEqual((t['invoices'], t['billed'], t['outstanding']), (1, '1180.00', '680.00'))

    def test_every_standing_invoice_counts_and_a_credited_one_does_not(self):
        self.invoice('100', report=False)                                   # the legacy inner join dropped this
        self.invoice('200', cancelled=True)                                 # reversed by a credit note
        self.assertEqual(self.get('invoices').json()['totals']['billed'], '100.00')

    def test_outstanding_is_money_owed_on_work_not_cancelled(self):
        self.invoice('100', remaining='40')
        self.invoice('100', remaining='0')
        self.invoice('100', remaining='30', status='Cancel')
        t = self.get('outstanding').json()['totals']
        self.assertEqual((t['invoices'], t['outstanding']), (1, '40.00'))

    def test_each_invoice_falls_into_exactly_one_age_band(self):
        """Regression: the legacy bands overlapped at 60 and 90 days."""
        for days in (60, 61, 90, 91, 121):
            self.invoice('100', remaining='10', days_ago=days)
        bands = {g['label']: g['values']['invoices'] for g in self.get('pending-recovery').json()['groups']}
        self.assertEqual(bands, {'0 to 60 days': 1, '61 to 90 days': 2, '91 to 120 days': 1, 'Over 120 days': 1})

    def test_the_six_invoice_groupings(self):
        a, b = m.Employee.objects.create(name='Asha'), m.Employee.objects.create(name='Vivek')
        self.invoice('100', assigned=a, asset_district='Bhopal', asset_state='Madhya Pradesh')
        self.invoice('300', assigned=b, asset_district='Itarsi', asset_state='Madhya Pradesh')
        for group in ('area', 'assessor', 'customer', 'client', 'service', 'division'):
            body = self.get('invoices', group=group).json()
            self.assertEqual(body['group'], group)
            self.assertEqual(sum(g['values']['invoices'] for g in body['groups']), 2, group)
        by_assessor = {g['label']: g['values']['billed'] for g in self.get('invoices', group='assessor').json()['groups']}
        self.assertEqual(by_assessor, {'Vivek': '300.00', 'Asha': '100.00'})
        self.assertEqual([g['label'] for g in self.get('invoices', group='area').json()['groups']][0], 'Itarsi, Madhya Pradesh')

    def test_the_date_filter_uses_the_invoice_date(self):
        self.invoice('100', days_ago=5)
        self.invoice('200', days_ago=40)
        since = (self.today - datetime.timedelta(days=10)).isoformat()
        self.assertEqual(self.get('invoices', **{'from': since}).json()['totals']['billed'], '100.00')
        self.assertEqual(self.get('invoices', **{'from': self.today.isoformat(), 'to': since}).status_code, 400)

    def test_the_records_behind_a_group(self):
        inv = self.invoice('100', remaining='25')
        group = self.get('outstanding', group='client').json()['groups'][0]['key']
        detail = self.get('outstanding', group='client', within=group).json()['detail']
        self.assertEqual((detail['total'], detail['rows'][0]['invoice'], detail['rows'][0]['outstanding']), (1, inv.number, '25.00'))

    def test_excel_export_has_a_summary_and_the_details(self):
        self.invoice('3850000', remaining='50000')
        response = self.client.get('/api/mis/outstanding/export/', {'group': 'client'})
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        wb = load_workbook(io.BytesIO(response.content))
        self.assertEqual(wb.sheetnames, ['Summary', 'Details'])
        summary = [c.value for row in wb['Summary'].iter_rows() for c in row if c.value is not None]
        self.assertIn('Omen Assessors: Outstanding', summary)
        self.assertIn(50000.0, summary)
        money_cell = next(c for row in wb['Details'].iter_rows(min_row=2) for c in row if c.value == 3850000.0)
        self.assertIn('##\\,##', money_cell.number_format)                    # Indian digit grouping

    def test_billing_by_liable_assessor_resolves_legacy_employee_ids(self):
        fa = m.Employee.objects.create(name='Ravi Liable')
        m.SourceIdentity.objects.create(batch=m.ImportBatch.objects.create(), environment='main', schema='uam_user',
                                        table='employee', key=source_key('57'), target_model='target_core.Employee', target_id=fa.pk)
        inv = self.invoice('500')
        m.BankReport.objects.filter(valuation=inv.valuation).update(legacy_liable_fa='57')
        groups = self.get('billing-by-assessor').json()['groups']
        self.assertEqual([(g['label'], g['values']['reports'], g['values']['billed']) for g in groups], [('Ravi Liable', 1, '500.00')])

    def test_a_restricted_user_sees_only_their_own_figures(self):
        me = m.Employee.objects.create(name='Assessor')
        self.invoice('100', assigned=me)
        self.invoice('900', assigned=m.Employee.objects.create(name='Other'))
        self.client.force_login(user_with_role('157', me))
        self.assertEqual(self.get('invoices').json()['totals']['billed'], '100.00')

    def test_the_catalogue_and_an_unknown_report(self):
        keys = [r['key'] for r in self.client.get('/api/mis/').json()['reports']]
        self.assertEqual(keys, ['invoices', 'outstanding', 'pending-recovery', 'total-amount', 'billing-by-client', 'billing-by-assessor',
                                'leads', 'lead-mis', 'valuations', 'doer-wise', 'service-rendered', 'cancelled-valuations',
                                'offline-billing', 'fa-targets'])
        self.assertEqual(self.get('no-such-report').status_code, 404)


class WorkReportTests(MisFixture, TestCase):
    def valuation(self, status, doer=None, reason_date=None, **kw):
        c = self.lead_with_case(**kw)
        return m.Valuation.objects.create(case=c, legacy_status=status, assessor=doer, valuation_date=self.today,
                                          status_reason_date=reason_date, status_reason='Double entry' if status == 'Cancel' else '')

    def test_leads_leave_out_deleted_ones_and_count_live_borrowers(self):
        """The legacy lead queries had no deleted-lead filter."""
        live = self.lead_with_case().lead
        m.Borrower.objects.create(lead=live, name='Second party', legacy_deleted='1')
        gone = self.lead_with_case().lead
        m.Lead.objects.filter(pk=gone.pk).update(legacy_deleted='1')
        t = self.get('leads').json()['totals']
        self.assertEqual((t['leads'], t['borrowers']), (1, 1))

    def test_lead_statuses_keep_their_natural_order(self):
        for status in ('4', '1', '2', '4'):
            m.Lead.objects.filter(pk=self.lead_with_case().lead_id).update(legacy_status=status)
        self.assertEqual([g['label'] for g in self.get('leads', group='status').json()['groups']], ['New lead', 'In process', 'Closed'])

    def test_lead_mis_counts_each_status_per_client(self):
        m.Lead.objects.filter(pk=self.lead_with_case().lead_id).update(legacy_status='2')
        t = self.get('lead-mis').json()['totals']
        self.assertEqual((t['leads'], t['in_process'], t['closed']), (1, 1, 0))

    def test_doer_wise_counts_each_status_and_runs_largest_first(self):
        busy, quiet = m.Employee.objects.create(name='Rahul'), m.Employee.objects.create(name='Aarti')
        for status in ('Submit', 'Submit', 'Revert'):
            self.valuation(status, doer=busy)
        self.valuation('Pending', doer=quiet)
        groups = self.get('doer-wise').json()['groups']
        self.assertEqual([(g['label'], g['values']['valuations'], g['values']['submitted'], g['values']['reverted']) for g in groups],
                         [('Rahul', 3, 2, 1), ('Aarti', 1, 0, 0)])

    def test_service_rendered_counts_each_invoice_once(self):
        """Regression: the legacy query counted 71,592 invoices for 11,926."""
        self.invoice('1180')
        t = self.get('service-rendered').json()['totals']
        self.assertEqual((t['invoices'], t['billed']), (1, '1180.00'))

    def test_every_cancelled_valuation_is_listed(self):
        """Regression: the legacy report showed one cancellation per client."""
        office = m.ClientOffice.objects.create(name='SBI Main')
        for _ in range(3):
            c = self.lead_with_case()
            m.Lead.objects.filter(pk=c.lead_id).update(office=office)
            m.Valuation.objects.create(case=c, legacy_status='Cancel', status_reason='Wrong entry', status_reason_date=self.today)
        body = self.get('cancelled-valuations').json()
        self.assertEqual((body['totals']['cancelled'], len(body['groups'])), (3, 1))
        detail = self.get('cancelled-valuations', within=body['groups'][0]['key']).json()['detail']
        self.assertEqual(detail['total'], 3)

    def test_cancellations_are_dated_by_when_they_were_cancelled(self):
        self.valuation('Cancel', reason_date=self.today - datetime.timedelta(days=40))
        self.valuation('Cancel', reason_date=self.today)
        since = (self.today - datetime.timedelta(days=7)).isoformat()
        self.assertEqual(self.get('cancelled-valuations', **{'from': since}).json()['totals']['cancelled'], 1)

    def test_recovery_runs_largest_debt_first(self):
        small, large = m.ClientOffice.objects.create(name='Small'), m.ClientOffice.objects.create(name='Large')
        for office, billed, owed in ((small, '10000', '100'), (large, '500', '400')):
            inv = self.invoice(billed, remaining=owed)
            m.Lead.objects.filter(pk=inv.valuation.case.lead_id).update(office=office)
        self.assertEqual([g['label'] for g in self.get('outstanding', group='client').json()['groups']][:2], ['Large', 'Small'])


class LastReportTests(MisFixture, TestCase):
    """Offline billing (10) and field assessors' targets (11)."""

    def offline(self, billed, advance='0', paid=(), days_ago=30, client='State Bank of India'):
        inv = m.Invoice.objects.create(origin='legacy_offline_invoice', number=f'I-OFF{m.Invoice.objects.count()}',
                                       issued_date=self.today - datetime.timedelta(days=days_ago), billed_amount=Decimal(billed),
                                       advance_amount=Decimal(advance), billed_to_name=client)
        for amount in paid:
            m.Receipt.objects.create(origin='legacy_offline_amount_received', invoice=inv, amount=Decimal(amount))
        return inv

    def target(self, employee, amount, cases=10, year=2026, source_id=None):
        t = m.EmployeeTarget.objects.create(employee=employee, financial_year=str(year), period_start=datetime.date(year, 4, 1),
                                            period_end=datetime.date(year + 1, 3, 31), target_cases=cases, target_amount=Decimal(amount),
                                            legacy_status='1')
        if source_id:
            m.SourceIdentity.objects.create(batch=m.ImportBatch.objects.create(), environment='main', schema='lead_mngt',
                                            table='employee_target', key=source_key(source_id), target_model='target_core.EmployeeTarget', target_id=t.pk)
        return t

    def liable(self, employee, billed, on):
        inv = self.invoice(billed)
        m.Invoice.objects.filter(pk=inv.pk).update(issued_date=on)
        m.Valuation.objects.filter(pk=inv.valuation_id).update(liable_employee=employee)
        return inv

    def test_offline_billing_shows_what_is_still_owed(self):
        """Regression: the legacy query grouped by a table it did not select, and always failed."""
        self.offline('1000', advance='200', paid=['300'])
        self.offline('500', paid=['500'], client='Canara Bank')
        t = self.get('offline-billing').json()['totals']
        self.assertEqual((t['invoices'], t['billed'], t['paid'], t['outstanding']), (2, '1500.00', '1000.00', '500.00'))
        self.assertEqual(self.get('offline-billing').json()['groups'][0]['label'], 'State Bank of India')   # largest debt first

    def test_offline_billing_is_hidden_from_a_restricted_user(self):
        self.offline('1000')
        self.client.force_login(user_with_role('157', m.Employee.objects.create(name='Assessor')))
        self.assertEqual(self.get('offline-billing').json()['totals']['invoices'], 0)

    def test_achievement_counts_invoices_within_the_financial_year(self):
        fa = m.Employee.objects.create(name='Noopur')
        self.target(fa, '100000', cases=4)
        self.liable(fa, '20000', datetime.date(2026, 5, 1))
        self.liable(fa, '30000', datetime.date(2027, 3, 31))
        self.liable(fa, '99999', datetime.date(2026, 3, 31))                 # the year before
        row = self.get('fa-targets').json()['groups'][0]['values']
        self.assertEqual((row['achieved_cases'], row['achieved_amount'], row['share']), (2, '50000.00', '50.00'))

    def test_a_duplicated_target_counts_achievement_once(self):
        """Regression: the legacy grouping joined both targets to the same invoices."""
        fa = m.Employee.objects.create(name='Noopur')
        self.target(fa, '100000', source_id='9')
        self.target(fa, '200000', source_id='12')                            # the more recent; 12 is newer than 9
        self.liable(fa, '50000', datetime.date(2026, 6, 1))
        body = self.get('fa-targets', group='employee').json()
        self.assertEqual((body['totals']['target_amount'], body['totals']['achieved_amount'], body['totals']['share']),
                         ('200000.00', '50000.00', '25.00'))
        detail = self.get('fa-targets', group='employee', within=body['groups'][0]['key']).json()['detail']
        self.assertEqual(detail['rows'][0]['superseded'], 1)

    def test_a_target_set_here_supersedes_an_imported_one(self):
        fa = m.Employee.objects.create(name='Noopur')
        self.target(fa, '100000', source_id='40')
        self.target(fa, '300000')
        self.assertEqual(self.get('fa-targets').json()['totals']['target_amount'], '300000.00')

    def test_percentages_are_worked_out_from_totals_not_added_up(self):
        a, b = m.Employee.objects.create(name='A'), m.Employee.objects.create(name='B')
        self.target(a, '100000')
        self.target(b, '300000')
        self.liable(a, '100000', datetime.date(2026, 6, 1))                   # A: 100%; B: 0%
        self.assertEqual(self.get('fa-targets').json()['totals']['share'], '25.00')   # not 100 + 0

    def test_the_bank_report_decides_who_is_liable(self):
        """Where the report and the valuation disagree, the report is the document of record."""
        on_report, on_valuation = m.Employee.objects.create(name='On report'), m.Employee.objects.create(name='On valuation')
        m.SourceIdentity.objects.create(batch=m.ImportBatch.objects.create(), environment='main', schema='uam_user',
                                        table='employee', key=source_key('82'), target_model='target_core.Employee', target_id=on_report.pk)
        inv = self.liable(on_valuation, '5151', datetime.date(2026, 6, 1))
        m.BankReport.objects.filter(valuation=inv.valuation).update(legacy_liable_fa='82')
        self.assertEqual([g['label'] for g in self.get('billing-by-assessor').json()['groups']], ['On report'])

    def test_percentages_export_as_percentages(self):
        fa = m.Employee.objects.create(name='Noopur')
        self.target(fa, '100000')
        wb = load_workbook(io.BytesIO(self.client.get('/api/mis/fa-targets/export/').content))
        formats = {c.number_format for row in wb['Summary'].iter_rows() for c in row}
        self.assertIn('0.0"%"', formats)
