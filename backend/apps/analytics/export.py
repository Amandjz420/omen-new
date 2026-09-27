"""Excel export for the MIS: a summary sheet and a detail sheet, as the legacy
exports offered, with the filters that produced them written at the top."""
import io
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Indian digit grouping in Excel: 1,23,45,678.00.
INDIAN = '[>=10000000]##\\,##\\,##\\,##0.00;[>=100000]##\\,##\\,##0.00;##,##0.00'
HEAD = PatternFill('solid', fgColor='357ABF')


def workbook(spec, f, group_label, groups, totals):
    wb = Workbook()
    summary = wb.active
    summary.title = 'Summary'
    period = ' to '.join(d.strftime('%d %b %Y') for d in (f['from'], f['to']) if d) or 'All dates'
    lines = [f"Omen Assessors: {spec['title']}", f"{spec['date_label']}: {period}",
             f"Generated {timezone.localtime():%d %b %Y, %I:%M %p}"]
    lines += [f'{k.title()}: {v}' for k, v in f.items() if v and k not in ('from', 'to')]
    for i, text in enumerate(lines, 1):
        summary.cell(row=i, column=1, value=text).font = Font(bold=(i == 1), size=13 if i == 1 else 10)
    start = len(lines) + 2
    measures = spec['measures']
    header = [group_label] + [x[1] for x in measures]
    for c, text in enumerate(header, 1):
        cell = summary.cell(row=start, column=c, value=text)
        cell.font, cell.fill = Font(bold=True, color='FFFFFF'), HEAD
    for r, g in enumerate(groups, start + 1):
        summary.cell(row=r, column=1, value=g['label'])
        for c, measure in enumerate(measures, 2):
            k, kind = measure[0], measure[2]
            v = g['values'][k]
            cell = summary.cell(row=r, column=c, value=float(v) if kind in ('money', 'percent') and v is not None else v)
            if kind == 'money':
                cell.number_format = INDIAN
            elif kind == 'percent':
                cell.number_format = '0.0"%"'
    last = start + len(groups) + 1
    summary.cell(row=last, column=1, value='Total').font = Font(bold=True)
    for c, measure in enumerate(measures, 2):
        k, kind = measure[0], measure[2]
        v = totals[k]
        cell = summary.cell(row=last, column=c, value=float(v) if kind in ('money', 'percent') and v is not None else v)
        cell.font = Font(bold=True)
        if kind == 'money':
            cell.number_format = INDIAN
        elif kind == 'percent':
            cell.number_format = '0.0"%"'
    widths(summary, [36] + [16] * len(measures))

    detail = wb.create_sheet('Details')
    cols = spec['columns']
    heads = [group_label] + [l for _, l, _ in cols]
    for c, text in enumerate(heads, 1):
        cell = detail.cell(row=1, column=c, value=text)
        cell.font, cell.fill = Font(bold=True, color='FFFFFF'), HEAD
    r = 2
    for g in groups:
        for row in g['rows']:
            detail.cell(row=r, column=1, value=g['label'])
            for c, (k, _, kind) in enumerate(cols, 2):
                v = row.get(k)
                cell = detail.cell(row=r, column=c, value=float(v) if isinstance(v, Decimal) else v)
                if kind == 'money':
                    cell.number_format = INDIAN
                elif kind == 'percent':
                    cell.number_format = '0.0"%"'
                elif kind == 'date' and v:
                    cell.number_format = 'DD-MMM-YYYY'
            r += 1
    detail.freeze_panes = 'A2'
    widths(detail, [30] + [16 if k in ('money', 'date') else 22 for _, _, k in cols])

    buffer = io.BytesIO()
    wb.save(buffer)
    name = f"Omen {spec['title']} {timezone.localdate():%Y-%m-%d}.xlsx"
    response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{name}"'
    return response


def widths(sheet, values):
    for i, w in enumerate(values, 1):
        sheet.column_dimensions[get_column_letter(i)].width = w
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical='top')
